from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


FUNC_RE = re.compile(r"(?<![A-Z0-9_.])([A-Z][A-Z0-9_.]*)\s*\(", re.I)
SHEET_REF_RE = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_\u4e00-\u9fff]+))!")


def normalize_formula(formula: str, coord: str) -> str:
    try:
        return Translator(formula, origin=coord).translate_formula("AZ100")
    except Exception:
        return formula


def nearby_labels(ws, row: int, col: int) -> list[str]:
    labels = []
    for c in range(max(1, col - 3), col):
        value = ws.cell(row, c).value
        if value not in (None, "") and not (isinstance(value, str) and value.startswith("=")):
            labels.append(str(value)[:80])
    return labels


def analyze(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=False)
    result = {"path": str(path), "sheets": [], "defined_names": [], "calc": {}}
    calc = wb.calculation
    result["calc"] = {
        "mode": getattr(calc, "calcMode", None),
        "iterate": getattr(calc, "iterate", None),
        "iterate_count": getattr(calc, "iterateCount", None),
        "iterate_delta": getattr(calc, "iterateDelta", None),
    }
    for dn in wb.defined_names.values():
        result["defined_names"].append({"name": dn.name, "attr_text": dn.attr_text})

    for ws in wb.worksheets:
        formulas = []
        funcs = Counter()
        external_refs = Counter()
        patterns = defaultdict(list)
        errors = []
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    formula = v
                    for fn in FUNC_RE.findall(formula.upper()):
                        funcs[fn] += 1
                    for m in SHEET_REF_RE.finditer(formula):
                        external_refs[m.group(1) or m.group(2)] += 1
                    pat = normalize_formula(formula, cell.coordinate)
                    patterns[pat].append(cell.coordinate)
                    formulas.append({
                        "cell": cell.coordinate,
                        "formula": formula,
                        "labels_left": nearby_labels(ws, cell.row, cell.column),
                    })
                    if any(e in formula.upper() for e in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!")):
                        errors.append({"cell": cell.coordinate, "formula": formula})
        result["sheets"].append({
            "title": ws.title,
            "state": ws.sheet_state,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "formula_count": len(formulas),
            "functions": funcs.most_common(),
            "sheet_refs": external_refs.most_common(),
            "formula_patterns": sorted(
                ({"normalized": k, "cells": v, "count": len(v)} for k, v in patterns.items()),
                key=lambda x: (-x["count"], x["cells"][0]),
            ),
            "formulas": formulas,
            "literal_formula_errors": errors,
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
        })
    return result


if __name__ == "__main__":
    reports = [analyze(Path(p)) for p in sys.argv[1:] if not p.startswith("--")]
    if "--compare" in sys.argv:
        out = []
        for report in reports:
            entry = {"file": Path(report["path"]).name, "sheets": []}
            for sh in report["sheets"]:
                coords = {f["cell"] for f in sh["formulas"]}
                constants = Counter()
                for f in sh["formulas"]:
                    for val in re.findall(r"(?<![A-Z0-9_])(\d+(?:\.\d+)?%?)(?![A-Z0-9_])", f["formula"], re.I):
                        if val not in {"0", "1", "10", "12"}:
                            constants[val] += 1
                out.append({
                    "file": entry["file"], "sheet": sh["title"], "size": sh["size"] if "size" in sh else f"{sh['max_row']}x{sh['max_col']}",
                    "formulas": sh["formula_count"], "functions": sh["functions"], "sheet_refs": sh["sheet_refs"],
                    "hardcoded_constants": constants.most_common(12),
                    "special": [f for f in sh["formulas"] if re.search(r"IRR|LOOKUP|#REF!|\[.*\]", f["formula"], re.I)],
                })
        pairwise = []
        for i in range(len(reports)):
            for j in range(i + 1, len(reports)):
                common_sheets = set(s["title"] for s in reports[i]["sheets"]) & set(s["title"] for s in reports[j]["sheets"])
                for name in common_sheets:
                    a = next(s for s in reports[i]["sheets"] if s["title"] == name)
                    b = next(s for s in reports[j]["sheets"] if s["title"] == name)
                    ca = Counter(p["normalized"] for p in a["formula_patterns"] for _ in range(p["count"]))
                    cb = Counter(p["normalized"] for p in b["formula_patterns"] for _ in range(p["count"]))
                    inter = sum((ca & cb).values()); union = sum((ca | cb).values())
                    pairwise.append({"a": Path(reports[i]["path"]).name, "b": Path(reports[j]["path"]).name, "sheet": name,
                                     "formula_similarity": round(inter / union, 4) if union else 1,
                                     "only_a_count": sum((ca - cb).values()), "only_b_count": sum((cb - ca).values()),
                                     "only_a_samples": list((ca - cb).keys())[:8], "only_b_samples": list((cb - ca).keys())[:8]})
        print(json.dumps({"summary": out, "pairwise": pairwise}, ensure_ascii=False, indent=2))
    elif "--summary" not in sys.argv:
        print(json.dumps(reports, ensure_ascii=False, indent=2))
    else:
        compact = []
        for report in reports:
            book = {"path": report["path"], "calc": report["calc"], "sheets": []}
            for sh in report["sheets"]:
                key = [
                    f for f in sh["formulas"]
                    if any(x in f["formula"].upper() for x in ("IRR(", "NPV(", "LOOKUP(", "RATE(", "25%", "0.25"))
                ]
                repeated = [p for p in sh["formula_patterns"] if p["count"] >= 3][:15]
                book["sheets"].append({
                    "title": sh["title"], "size": f"{sh['max_row']}x{sh['max_col']}",
                    "formula_count": sh["formula_count"], "functions": sh["functions"],
                    "sheet_refs": sh["sheet_refs"], "repeated_patterns": repeated,
                    "key_formulas": key[:30], "literal_formula_errors": sh["literal_formula_errors"],
                })
            compact.append(book)
        print(json.dumps(compact, ensure_ascii=False, indent=2))
