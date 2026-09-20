from pathlib import Path
import json, sys
from openpyxl import load_workbook

ranges = {
    "政企项目投资效益评估表": ["A53:L72", "A7:W18", "A49:W51"],
}

def read_rect(ws, ref):
    out=[]
    for row in ws[ref]:
        vals=[]
        for c in row:
            if c.value not in (None, ""):
                vals.append([c.coordinate, c.value])
        if vals: out.append(vals)
    return out

ans=[]
for arg in (x for x in sys.argv[1:] if not x.startswith("--")):
    p=Path(arg); wf=load_workbook(p,data_only=False); wv=load_workbook(p,data_only=True)
    item={"file":p.name,"ranges":{},"cached_errors":[]}
    for ws in wf.worksheets:
        for row in wv[ws.title].iter_rows():
            for c in row:
                if isinstance(c.value,str) and c.value.startswith("#"):
                    item["cached_errors"].append([ws.title,c.coordinate,c.value])
        if ws.title in ranges:
            item["ranges"][ws.title]={ref:read_rect(ws,ref) for ref in ranges[ws.title]}
    ans.append(item)
if "--errors" in sys.argv:
    print(json.dumps([{"file":x["file"],"cached_errors":x["cached_errors"]} for x in ans],ensure_ascii=False,indent=2))
else:
    print(json.dumps(ans,ensure_ascii=False,indent=2))
